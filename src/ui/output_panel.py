# src/ui/output_panel.py
# Right panel that displays OCR results in both raw text and rendered markdown.

import os
import re
import html
import markdown
import pypandoc
from PySide6.QtWidgets import QWidget, QVBoxLayout, QHBoxLayout, QTabWidget, QPushButton, QLabel, QTextEdit, QMenu, QFileDialog, QMessageBox
from PySide6.QtWebEngineWidgets import QWebEngineView
from PySide6.QtWebEngineCore import QWebEnginePage
from PySide6.QtCore import Qt, Slot, QUrl
from PySide6.QtGui import QTextCursor


# CSS styling for the fancy (rendered) output view
BROWSER_STYLE = """
<style>
body { background-color: #FFFFFF; color: #000000; font-family: "Times New Roman", serif; font-size: 13pt; line-height: 1.6; }
a { color: #0000EE; text-decoration: none; }
a:hover { text-decoration: underline; }
code { background-color: #F0F0F0; padding: 2px 5px; border-radius: 3px; font-family: Consolas, monospace; }
pre { background-color: #F0F0F0; padding: 10px; border-radius: 5px; overflow-x: auto; border: 1px solid #CCC; }
pre code { background-color: transparent; padding: 0; }
h1, h2, h3, h4, h5, h6 { color: #000000; margin-top: 1.5em; margin-bottom: 0.5em; }
h1 { border-bottom: 1px solid #CCCCCC; padding-bottom: 0.3em; }
blockquote { border-left: 4px solid #CCCCCC; margin-left: 0; padding-left: 15px; color: #666666; font-style: italic; }
ul, ol { padding-left: 25px; }
li { margin-bottom: 0.3em; }
table { border-collapse: collapse; width: 100%; margin-bottom: 1em; }
th, td { border: 1px solid #000000; padding: 8px; text-align: left; }
th { background-color: #F2F2F2; font-weight: bold; }
tr:nth-child(even) { background-color: #FAFAFA; }
img { max-width: 100%; max-height: 100%; height: auto; width: auto; display: block; margin: 10px auto; }
</style>
"""


def balance_latex_delimiters(latex):
    """
    Fix unbalanced \\left and \\right delimiters in LaTeX.
    LaTeX requires matching \\left and \\right pairs. DeepSeek-OCR sometime
    outputs unbalanced delimiters, which cause MathJax to fail when rendering.
    This function:
    1. Removes orphan \\right commands (no matching \\left)
    2. Adds \\right. (invisible delimiter) for unmatched \\left
    """
    # Find all \left and \right commands with their positions
    commands = []
    for m in re.finditer(r'(\\left|\\right)', latex):
        commands.append((m.start(), m.end(), m.group(0)))

    to_remove = set()
    stack = 0

    # Track depth: \left increases, \right decreases
    for start, end, cmd in commands:
        if cmd == r'\left':
            stack += 1
        else:
            if stack > 0:
                stack -= 1
            else:
                # Unmatched \right - mark for removal
                to_remove.add(start)

    # Rebuild string, skipping marked positions
    fixed_latex = ""
    last_idx = 0
    for start, end, cmd in commands:
        if start in to_remove:
            fixed_latex += latex[last_idx:start]
            last_idx = end

    fixed_latex += latex[last_idx:]

    # Add invisible delimiters for unmatched \left
    if stack > 0:
        fixed_latex += (r" \right." * stack)

    return fixed_latex


# ==================== Tab 2: Fancy Output ====================
class FancyOutput(QWebEngineView):
    # WebView that renders markdown with MathJax LaTeX support.
    def __init__(self, parent=None):
        super().__init__(parent)
        self.page().setBackgroundColor(Qt.white)

    def _convert_to_html(self, raw_md):
        # 1. Extract and protect LaTeX blocks from markdown processing
        # Markdown would destroy LaTeX syntax, so we replace with placeholders
        math_blocks = []

        def replace_math(match):
            block = match.group(0)
            block = balance_latex_delimiters(block)
            math_blocks.append(block)
            return f"MATHJAXBLOCKPLACEHOLDER{len(math_blocks)-1}END"

        # Match both display (\[...\]) and inline (\(...\)) math
        pattern = re.compile(r'(\\\[.*?\\\])|(\\\(.*?\\\))', re.DOTALL)
        processed_md = pattern.sub(replace_math, raw_md)

        # 2. Convert markdown to HTML
        html_content = markdown.markdown(processed_md, extensions=['tables', 'fenced_code', 'nl2br'])

        # 3. Restore LaTeX blocks, escaped for HTML safety
        for i, block in enumerate(math_blocks):
            safe_block = html.escape(block)
            html_content = html_content.replace(f"MATHJAXBLOCKPLACEHOLDER{i}END", safe_block)

        return html_content

    def set_markdown(self, md_content):
        # Convert markdown to HTML and render with MathJax.
        if not md_content:
            self.setHtml("")
            return

        try:
            html_content = self._convert_to_html(md_content)

            # 4. Setup MathJax
            base_path = os.path.dirname(os.path.abspath(__file__))
            project_path = os.path.dirname(base_path)
            node_path = os.path.join(project_path, "res", "node")

            mathjax_path = os.path.join(node_path, "mathjax", "tex-mml-svg.js")
            mathjax_url = QUrl.fromLocalFile(mathjax_path).toString()
            mathjax_dir_url = QUrl.fromLocalFile(os.path.dirname(mathjax_path)).toString()
            node_path_url = QUrl.fromLocalFile(node_path).toString()

            # Explicitly map the newcm font to ensure local loading
            newcm_path = os.path.join(node_path, "@mathjax", "mathjax-newcm-font")
            newcm_url = QUrl.fromLocalFile(newcm_path).toString()

            # STUPID Hack: catch and replace cdn link with local path (setting paths doesn't work)
            mathjax_script = f"""
            <script>
            MathJax = {{
              loader: {{
                paths: {{
                  mathjax: '{mathjax_dir_url}',
                  npm: '{node_path_url}',
                  'mathjax-newcm-font': '{newcm_url}'
                }},
                pathFilters: [
                  [(data) => {{
                    var cdn = 'https://cdn.jsdelivr.net/npm/';
                    if (data.name.indexOf(cdn) === 0) {{
                        data.name = data.name.replace(cdn, '{node_path_url}/');
                    }}
                    return true;
                  }}, 25]
                ]
              }},
              tex: {{
                inlineMath: [['\\\\(', '\\\\)']],
                displayMath: [['\\\\[', '\\\\]']],
                processEscapes: true
              }},
              svg: {{
                fontCache: 'global'
              }},
              options: {{
                ignoreHtmlClass: 'tex2jax_ignore',
                processHtmlClass: 'tex2jax_process',
                menuOptions: {{
                  settings: {{
                    assistiveMml: true,
                    enrich: false
                  }}
                }}
              }}
            }};
            </script>
            <script id="MathJax-script" async src="{mathjax_url}"></script>
            """

            full_html = f"<html><head>{BROWSER_STYLE}{mathjax_script}</head><body>{html_content}</body></html>"
            self.setHtml(full_html, QUrl.fromLocalFile(project_path))
        except Exception as e:
            print(f"Markdown render error: {e}")

    def copy_content(self):
        self.setFocus()
        page = self.page()
        page.triggerAction(QWebEnginePage.WebAction.SelectAll)
        page.triggerAction(QWebEnginePage.WebAction.Copy)

    def contextMenuEvent(self, event):
        menu = QMenu(self)
        menu.addAction(self.page().action(QWebEnginePage.WebAction.Copy))
        menu.addAction(self.page().action(QWebEnginePage.WebAction.SelectAll))
        menu.exec(event.globalPos())


# ==================== Main Widget ====================
class OutputPanel(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout(self)
        self.t = {}  # Translation dictionary

        # === Tabs ===
        self.tabs = QTabWidget()

        # Tab 1: Raw Output
        self.tab_raw = QWidget()
        raw_layout = QVBoxLayout(self.tab_raw)
        raw_layout.setContentsMargins(0, 0, 0, 0)
        self.text_output = QTextEdit()
        self.text_output.setReadOnly(True)
        raw_layout.addWidget(self.text_output)

        # Tab 2: Fancy Output
        self.web_view = FancyOutput()

        self.tabs.addTab(self.tab_raw, "")
        self.tabs.addTab(self.web_view, "")
        self.tabs.setTabEnabled(1, False)  # Disabled until processing completes
        self.tabs.currentChanged.connect(self._update_copy_button_text)

        # === Bottom Controls ===
        self.lbl_proofread = QLabel()
        self.lbl_proofread.setObjectName("lbl_proofread")
        self.lbl_proofread.setAlignment(Qt.AlignCenter)

        self.btn_copy = QPushButton()
        self.btn_copy.clicked.connect(self.copy_output)

        self.btn_export = QPushButton()
        self.btn_export.clicked.connect(self.export_to_word)

        btn_layout = QHBoxLayout()
        btn_layout.addWidget(self.btn_copy)
        btn_layout.addWidget(self.btn_export)

        self.layout.addWidget(self.tabs)
        self.layout.addWidget(self.lbl_proofread)
        self.layout.addLayout(btn_layout)

    # ==================== Language ====================
    def update_language(self, t):
        self.t = t
        self.lbl_proofread.setText(t["lbl_proofread"])
        self.tabs.setTabText(0, t["tab_raw"])
        self.tabs.setTabText(1, t["tab_fancy"])
        self.btn_export.setText(t.get("btn_export_word", "Export to Word"))
        self._update_copy_button_text()

    def set_processing_state(self, is_processing):
        self.btn_copy.setEnabled(not is_processing)
        self.btn_export.setEnabled(not is_processing)

    def _update_copy_button_text(self):
        if not self.t or "btn_copy" not in self.t:
            return
        current_tab_text = self.tabs.tabText(self.tabs.currentIndex())
        self.btn_copy.setText(self.t["btn_copy"].format(current_tab_text))

    # ==================== Tab 1: Raw Output ====================
    @Slot(str)
    def append_text(self, text):
        # Append streaming text to raw output. Switches to raw tab if on fancy.
        if self.tabs.currentIndex() != 0:
             self.tabs.setCurrentIndex(0)

        if self.tabs.isTabEnabled(1):
            self.tabs.setTabEnabled(1, False)

        # Move cursor to end, insert text, keep cursor at end (auto-scroll)
        self.text_output.moveCursor(QTextCursor.End)
        self.text_output.insertPlainText(text)
        self.text_output.moveCursor(QTextCursor.End)

    # ==================== Tab 2: Fancy Output ====================
    def render_fancy_output(self):
        # Convert raw text to rendered markdown.
        raw_md = self.text_output.toPlainText()
        if not raw_md.strip():
            return

        self.web_view.set_markdown(raw_md)
        self.tabs.setTabEnabled(1, True)
        self.tabs.setCurrentIndex(1)

    # ==================== Utility ====================
    def clear(self):
        self.text_output.clear()
        self.web_view.set_markdown("")
        self.tabs.setTabEnabled(1, False)
        self.tabs.setCurrentIndex(0)

    def copy_output(self):
        if self.tabs.currentIndex() == 0:
            self.text_output.selectAll()
            self.text_output.copy()
        else:
            self.web_view.copy_content()

    def export_to_word(self):
        raw_md = self.text_output.toPlainText()
        if not raw_md.strip():
            return

        filename, _ = QFileDialog.getSaveFileName(
            self,
            "Save Word Document",
            "",
            "Word Document (*.docx)"
        )

        if filename:
            try:
                import docx
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn

                # Convert to HTML first to preserve HTML tables and image paths
                # Then use Pandoc's HTML reader with math extensions to create DOCX
                html_content = self.web_view._convert_to_html(raw_md)
                pypandoc.convert_text(html_content, 'docx', format='html+tex_math_single_backslash', outputfile=filename)

                # Add borders to all tables via XML to avoid missing style errors
                doc = docx.Document(filename)
                for table in doc.tables:
                    tbl = table._tbl
                    tblPr = tbl.tblPr
                    if tblPr is None:
                        tblPr = OxmlElement('w:tblPr')
                        tbl.insert(0, tblPr)

                    tblBorders = tblPr.find(qn('w:tblBorders'))
                    if tblBorders is not None:
                        tblPr.remove(tblBorders)

                    tblBorders = OxmlElement('w:tblBorders')
                    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                        border = OxmlElement(f'w:{border_name}')
                        border.set(qn('w:val'), 'single')
                        border.set(qn('w:sz'), '4')  # 1/2 pt border
                        border.set(qn('w:space'), '0')
                        border.set(qn('w:color'), 'auto')
                        tblBorders.append(border)

                    tblPr.append(tblBorders)

                doc.save(filename)

            except Exception as e:
                QMessageBox.critical(self, "Export Error", f"Failed to export to Word:\n{e}")
